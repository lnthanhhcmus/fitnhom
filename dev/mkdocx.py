#!/usr/bin/env python3
"""Dựng bản Word (.docx, khổ ngang) từ một tệp Markdown của dự án, tùy chọn kèm thư mục ảnh.

Dùng:  python3 dev/mkdocx.py tai-lieu.md ra.docx [thu-muc-anh]
- Nếu có thư mục ảnh: mọi *.png trong đó (theo tên) được nối vào cuối dưới tiêu đề "# N. Ảnh minh họa" (N = số mục cuối + 1)
  — hoặc thay thế mục đó nếu tệp md đã có tiêu đề "# N. Ảnh minh họa".
- Cần pandoc và python-docx (pip install python-docx). Bảng được kẻ viền, chữ 9,5 pt, chiều rộng cột theo nội dung.
"""
import glob
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

md, out = sys.argv[1], sys.argv[2]
imgdir = sys.argv[3] if len(sys.argv) > 3 else None
src = open(md, encoding="utf-8").read()
m = re.search(r"^# (\d+)\. Ảnh minh họa", src, re.M)
body = src[: m.start()] if m else src
if imgdir and glob.glob(os.path.join(imgdir, "*.png")):
    n = m.group(1) if m else str(max([int(x) for x in re.findall(r"^# (\d+)\.", body, re.M)] or [0]) + 1)
    body = body.rstrip() + f"\n\n# {n}. Ảnh minh họa\n\nẢnh chụp màn hình kèm theo tài liệu.\n\n"
    for i, f in enumerate(sorted(glob.glob(os.path.join(imgdir, "*.png"))), 1):
        body += "![Ảnh %02d — %s](%s){width=15cm}\n\n" % (i, os.path.basename(f)[:-4], f)
tmp = out + ".md"
open(tmp, "w", encoding="utf-8").write(body)
subprocess.run(["pandoc", tmp, "-o", out, "--resource-path", os.path.dirname(os.path.abspath(md))], check=True)
d = Document(out)
for s in d.sections:
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.6)
usable = 29.7 - 3.6
for t in d.tables:
    ncol = len(t.columns)
    lens = [max(len(r.cells[i].text) for r in t.rows) for i in range(ncol)]
    tot = sum(min(x, 120) for x in lens) or 1
    w = [max(2.5, usable * min(x, 120) / tot) for x in lens]
    k = usable / sum(w)
    w = [x * k for x in w]
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement("w:" + e)
        x.set(qn("w:val"), "single")
        x.set(qn("w:sz"), "4")
        x.set(qn("w:color"), "999999")
        b.append(x)
    tblpr.append(b)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(w[i])
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
d.save(out)
os.remove(tmp)
print("ok", out)
